import asyncio
import base64
import io
import websockets
from PIL import Image
from docx import Document
from pptx import Presentation
import openpyxl


class WebSocketServer:
    def __init__(self, host, port):
        self.host = host
        self.port = port
        self.fileName = None
        self.finish = False

    async def receive_file_chunks(self, websocket):
        file_data = ""
        print("waiting chunks")
        chunk = await websocket.recv()

        if "fileName" in chunk:
            # get filename
            self.fileName = chunk.split(":")[1]

            chunk = await websocket.recv()
            while chunk is not None:
                # wait recv all chunks
                file_data += chunk
                try:
                    chunk = await websocket.recv()
                except websockets.exceptions.ConnectionClosedOK:
                    if len(file_data):
                        file_content = base64.b64decode(file_data)
                        fileNameExtension = self.fileName.split(".")[-1]

                        if fileNameExtension == "docx" or \
                                fileNameExtension == "doc":
                            document = Document()
                            document.save(self.fileName)
                            with open(self.fileName, "wb") as f:
                                f.write(file_content)

                        elif fileNameExtension == "pptx" or \
                                fileNameExtension == "ppt":
                            presentation = Presentation()
                            presentation.save(self.fileName)
                            with open(self.fileName, "wb") as f:
                                f.write(file_content)

                        elif fileNameExtension == "xlsx" or \
                                fileNameExtension == "xls":
                            workbook = openpyxl.Workbook()
                            workbook.save(self.fileName)
                            with open(self.fileName, "wb") as f:
                                f.write(file_content)

                        elif fileNameExtension == "jpg" or \
                                fileNameExtension == "jpeg" or \
                                fileNameExtension == "png":
                            pil_image = Image.open(io.BytesIO(file_content)).convert('RGB')
                            pil_image.save(self.fileName)

                        elif fileNameExtension == "webm" or \
                                fileNameExtension == "mkv" or \
                                fileNameExtension == "flv" or \
                                fileNameExtension == "mp4" or \
                                fileNameExtension == "gif" or \
                                fileNameExtension == "vob" or \
                                fileNameExtension == "ogv" or \
                                fileNameExtension == "ogg" or \
                                fileNameExtension == "avi" or \
                                fileNameExtension == "mov":
                            with open(self.fileName, "wb") as f:
                                f.write(file_content)

                        else:
                            with open(self.fileName, 'wb') as f:
                                f.write(file_content)
                    chunk = None

            print('File received successfully!')
            self.finish = True

    async def run(self):
        async with websockets.serve(self.receive_file_chunks, self.host, self.port):
            await asyncio.Future()  # run forever
